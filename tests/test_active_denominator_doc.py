"""The active-market denominator the documents state is the one the workbook gives (review of 2 October 2026, E-4 / R11-01).

The README and docs/data-audit-results.md called the active-market denominator "the 1,040 SFCR count the spreadsheet's own
note directs use of". Neither half held. The workbook's SFCR column totals 972. 1,040 was the analysis's count before it put
Syndicate 33 back on the 2020-2024 lists (analysis 8f0e6b8, 1 October 2026), and the paper has used 1,045 since then: the SFCR
column for 2014-2019 (572) and Lloyd's official lists of active syndicates for 2020-2024 (473). The analysis holds those lists
as the per-year sheets of this repository's workbook (its src/test_market_active.py holds the two together), and its
annual-report counts for 2014-2019 are the workbook's SFCR column.

So the counts are computed here from the workbook, and every document that states the denominator is held to them.

Run:  python -m pytest tests/test_active_denominator_doc.py -q
"""
import functools
import re
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
WORKBOOK = ROOT / "syndicate_reports" / "Lloyds_Syndicates_2014_2024.xlsx"
#: the documents that state the active-market denominator
DOCS = ("README.md", "docs/data-audit-results.md")
#: the years whose count is the workbook's SFCR column; the official lists cover the rest
SFCR_YEARS = range(2014, 2020)
LIST_YEARS = range(2020, 2025)


@functools.lru_cache(maxsize=None)
def _counts():
    wb = openpyxl.load_workbook(WORKBOOK, read_only=True, data_only=True)
    sfcr = {}
    for row in wb["Overview"].iter_rows(values_only=True):
        cells = [c for c in row if c is not None]
        if len(cells) >= 3 and isinstance(cells[0], int) and 2014 <= cells[0] <= 2024:
            sfcr[cells[0]] = cells[1]
    lists = {y: {int(r[0]) for r in wb[str(y)].iter_rows(min_row=2, values_only=True) if isinstance(r[0], (int, float))}
             for y in LIST_YEARS}
    rows = sum(1 for r in wb["All Data"].iter_rows(min_row=2, values_only=True) if r[0] is not None)
    wb.close()
    assert sorted(sfcr) == list(range(2014, 2025)), "the Overview sheet's SFCR column is not where this test reads it"
    early = sum(sfcr[y] for y in SFCR_YEARS)
    late = sum(len(lists[y]) for y in LIST_YEARS)
    return {"early": early, "late": late, "total": early + late, "sfcr_all": sum(sfcr.values()), "rows": rows}


def _fmt(n):
    return "{:,}".format(n)


def _flat(rel):
    return " ".join((ROOT / rel).read_text(encoding="utf-8").split())


def _units(rel):
    """The document's sentences: each table row is one unit, and each other paragraph is split into its sentences."""
    out = []
    for para in re.split(r"\n\s*\n", (ROOT / rel).read_text(encoding="utf-8")):
        rows = [line for line in para.splitlines() if line.lstrip().startswith("|")]
        out += rows
        prose = " ".join(" ".join(line for line in para.splitlines() if not line.lstrip().startswith("|")).split())
        out += [s for s in re.split(r"(?<=[.;])\s+(?=[A-Z(*])", prose) if s]
    return out


def test_the_workbook_gives_the_papers_denominator():
    """The counts the documents are held to, as measured on 2 October 2026 (a workbook change that moves them fails here
    first, and the documents then follow it)."""
    c = _counts()
    assert (c["early"], c["late"], c["total"], c["sfcr_all"], c["rows"]) == (572, 473, 1045, 972, 1125), c


def test_each_document_states_the_denominator_the_workbook_gives():
    c = _counts()
    for rel in DOCS:
        about = [s for s in _units(rel) if re.search(r"active[- ]market denominator|active syndicate-years", s)]
        assert about, (rel, "no sentence states the active-market denominator")
        assert any(_fmt(c["total"]) in s for s in about), (rel, "does not state %s" % _fmt(c["total"]), about)
        # a sentence about the SFCR column, or about active syndicates in any words, gives no other total either
        near = [s for s in _units(rel) if re.search(r"SFCR|\bactive\b", s)]
        for s in near:
            for n in re.findall(r"\b\d{1,3}(?:,\d{3})+\b", s):
                assert n in (_fmt(c["total"]), _fmt(c["rows"])), (rel, "a denominator the workbook does not give", n, s)


def test_the_parts_and_the_sfcr_total_are_the_workbooks():
    c = _counts()
    for rel in DOCS:
        text = _flat(rel)
        m = re.search(r"(\d+) for 2014[-–]2019", text)
        assert m and int(m.group(1)) == c["early"], (rel, "the 2014-2019 part", m and m.group(0))
        m = re.search(r"(\d+) for 2020[-–]2024", text)
        assert m and int(m.group(1)) == c["late"], (rel, "the 2020-2024 part", m and m.group(0))
        m = re.search(r"SFCR column, which totals (\d[\d,]*)", text)
        assert m and m.group(1) == _fmt(c["sfcr_all"]), (rel, "the SFCR column's total", m and m.group(0))


def test_no_document_calls_the_denominator_the_sfcr_count():
    """The SFCR column totals 972 and gives the denominator for 2014-2019 only; "the SFCR count" named a number it is not."""
    for rel in DOCS:
        text = _flat(rel)
        assert not re.search(r"SFCR count", text), (rel, re.search(r".{80}SFCR count.{40}", text).group(0))
