"""The download ledger, the workbook and the corpus, and what the documents say about the filings the ledger does not hold
(review of 2 October 2026, P-28 / R9-03 / R5-F7).

docs/data-audit-results.md said that 33 corpus filings "were obtained in an earlier collection pass and are flagged
already_present rather than logged as downloads", a "ledger-completeness gap of 3.1%". The ledger,
syndicate_reports/download_status.json, has one row for each of the workbook's 1,125 rows and none for the 33, and never had
one: already_present is the detail of 588 rows it records as downloaded. The 33 are not rows of the workbook.

These tests hold that account to the files: the corpus is the ledger's downloaded rows plus the filings the audit page lists,
the counts the page and the README state are the files' own, and no document repeats the retired account. The author decided
on 2 October 2026 that the ledger gains no rows for the 33: they are listed apart, with their source addresses and file
fingerprints, in syndicate_reports/download_addendum.json, which tests/test_download_addendum.py holds to the corpus and the
ledger.

Run:  python -m pytest tests/test_download_ledger.py -q
"""
import functools
import json
import re
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
LEDGER = ROOT / "syndicate_reports" / "download_status.json"
WORKBOOK = ROOT / "syndicate_reports" / "Lloyds_Syndicates_2014_2024.xlsx"
REGISTER = ROOT / "pdf_extraction" / "audit" / "runoff_corpus_register.json"
AUDIT_PAGE = ROOT / "docs" / "data-audit-results.md"
README = ROOT / "README.md"
DOWNLOADED = "report downloaded"


@functools.lru_cache(maxsize=None)
def _ledger():
    return json.loads(LEDGER.read_text(encoding="utf-8"))


@functools.lru_cache(maxsize=None)
def _corpus():
    return frozenset(p.stem.replace("syndicate_", "") for p in (ROOT / "pdf_extraction").glob("syndicate_*.json")
                     if re.fullmatch(r"syndicate_\d+_\d{4}", p.stem))


@functools.lru_cache(maxsize=None)
def _workbook_rows():
    wb = openpyxl.load_workbook(WORKBOOK, read_only=True, data_only=True)
    rows = ["%d_%d" % (int(r[1]), int(r[0])) for r in wb["All Data"].iter_rows(min_row=2, values_only=True)
            if r[0] is not None and r[1] is not None]
    wb.close()
    return tuple(rows)


def _flat(path):
    return " ".join(path.read_text(encoding="utf-8").split())


def _off_ledger_item():
    """The audit page's bullet about the filings the ledger does not hold."""
    text = AUDIT_PAGE.read_text(encoding="utf-8")
    m = re.search(r"^- The (\d+) are not rows of the workbook.*?(?=^- |^#|\Z)", text, re.M | re.S)
    assert m, "docs/data-audit-results.md has no item on the filings the ledger does not hold"
    return int(m.group(1)), " ".join(m.group(0).split())


def _listed(item):
    m = re.search(r"They are (.*?)\.(?: |$)", item)
    assert m, "the item does not list the filings"
    return {"%s_%s" % (s, y) for s, y in re.findall(r"\b(\d{2,4})/(20\d\d)\b", m.group(1))}


def test_the_ledger_is_the_workbook():
    """One ledger row per workbook row, and no other: the downloader writes a row only for a row of the workbook."""
    rows = _workbook_rows()
    assert len(rows) == len(set(rows)) == len(_ledger()), (len(rows), len(_ledger()))
    assert set(rows) == set(_ledger())


def test_the_corpus_is_the_downloaded_rows_and_the_filings_the_audit_page_lists():
    n, item = _off_ledger_item()
    listed = _listed(item)
    downloaded = {k for k, v in _ledger().items() if v["status"] == DOWNLOADED}
    assert downloaded <= _corpus(), sorted(downloaded - _corpus())
    off = _corpus() - set(_ledger())
    assert listed == off, {"listed, not off the ledger": sorted(listed - off), "off the ledger, not listed": sorted(off - listed)}
    assert n == len(off), (n, len(off))
    assert not off & set(_workbook_rows()), "a listed filing is a row of the workbook"
    assert not (_corpus() & {k for k, v in _ledger().items() if v["status"] != DOWNLOADED}), "an unavailable row has a record"


def test_the_counts_the_audit_page_states_are_the_files():
    _, item = _off_ledger_item()
    statuses = [v["status"] for v in _ledger().values()]
    already = sum(1 for v in _ledger().values() if v["status"] == DOWNLOADED and v.get("detail") == "already_present")
    m = re.search(r"`already_present` is the detail of (\d+) rows", item)
    assert m and int(m.group(1)) == already, (m and m.group(0), already)
    m = re.search(r"reconciles the ([\d,]+) workbook rows \(([\d,]+) downloaded\)", item)
    assert m, "the item does not say what the coverage report reconciles"
    assert (int(m.group(1).replace(",", "")), int(m.group(2).replace(",", ""))) == (len(statuses), statuses.count(DOWNLOADED))
    # the register's reading of the listed filings
    reg = json.loads(REGISTER.read_text(encoding="utf-8"))
    category = {r["stem"].replace("syndicate_", ""): r["category"] for r in reg["records"]}
    listed = _listed(item)
    counts = {c: sum(1 for s in listed if category.get(s) == c) for c in ("WHOLE", "PART", "NOTCOUNT")}
    m = re.search(r"(\d+) of them are run-off years \((\d+) WHOLE, (\d+) PART", item)
    assert m and tuple(int(x) for x in m.groups()) == (counts["WHOLE"] + counts["PART"], counts["WHOLE"], counts["PART"]), (
        m and m.group(0), counts)
    reviewed = {r["stem"].replace("syndicate_", "") for r in reg["reviewed_not_run_off"]}
    expected = {
        "are NOTCOUNT": {s for s in listed if category.get(s) == "NOTCOUNT"},
        "was reviewed and states no run-off": listed & reviewed,
        "have no entry": {s for s in listed if s not in category and s not in reviewed},
    }
    for phrase, stems in expected.items():
        m = re.search(r"((?:\d{2,4}/20\d\d(?:,? and |, )?)+) " + re.escape(phrase), item)
        named = {"%s_%s" % (a, b) for a, b in re.findall(r"(\d{2,4})/(20\d\d)", m.group(1))} if m else set()
        assert named == stems, (phrase, sorted(named), sorted(stems))


def test_the_urls_the_audit_page_says_are_committed_are_there():
    """The page says the workbook and the ledger hold no source URL for the listed filings, and that a Lloyd's URL for some of
    them is in the market commentary's discovered sources. Both halves are held to the files."""
    _, item = _off_ledger_item()
    listed = _listed(item)
    m = re.search(r"A Lloyd's URL for (\w+) of them, (.*?), appears? in `([^`]+)`", item)
    assert m, "the item does not say where a URL for any of them is"
    named = {"%s_%s" % (a, b) for a, b in re.findall(r"(\d{2,4})/(20\d\d)", m.group(2))}
    assert named <= listed and len(named) == {"one": 1, "two": 2, "three": 3, "four": 4}.get(m.group(1), -1), m.group(0)
    urls = set(re.findall(r"https?://[^\"'\s]*lloyds[^\"'\s]*", (ROOT / m.group(3)).read_text(encoding="utf-8")))

    def has_url(stem):
        syn, year = stem.split("_")
        pat = re.compile(r"(?<!\d)0*%s(?!\d)" % syn)
        return any(year in u and pat.search(u) and re.search(r"\.(pdf|html?)\b", u, re.I) for u in urls)

    assert {s for s in listed if has_url(s)} == named
    wb = openpyxl.load_workbook(WORKBOOK, read_only=True, data_only=True)
    workbook_urls = " ".join(str(r[3]) for r in wb["All Data"].iter_rows(min_row=2, values_only=True))
    wb.close()
    ledger_urls = " ".join(str(v.get("source_url", "")) for v in _ledger().values())
    for s in listed:
        syn, year = s.split("_")
        for where, text in (("workbook", workbook_urls), ("ledger", ledger_urls)):
            assert not re.search(r"%s[-_/]0*%s\b|\b0*%s[-_/]%s" % (year, syn, syn, year), text), (s, where)


def test_the_readme_download_step_states_the_files_counts():
    text = _flat(README)
    start = text.index("#### 1. Download syndicate reports")
    step = text[start:text.index("####", start + 10)]
    statuses = [v["status"] for v in _ledger().values()]
    off = len(_corpus() - set(_ledger()))
    assert "each of the {:,} rows".format(len(_workbook_rows())) in step
    m = re.search(r"This route gives ([\d,]+) of the corpus's ([\d,]+) filings; the ledger records the other (\d+) rows as "
                  r"unavailable", step)
    assert m, "the step does not say what the route gives"
    assert tuple(int(x.replace(",", "")) for x in m.groups()) == (
        statuses.count(DOWNLOADED), len(_corpus()), len(statuses) - statuses.count(DOWNLOADED)), m.group(0)
    m = re.search(r"The remaining (\d+) filings are not rows of the workbook", step)
    assert m and int(m.group(1)) == off, (m and m.group(0), off)
    assert "scripts/download_from_xlsx.py" in step


def test_no_document_repeats_the_retired_account_of_the_off_ledger_filings():
    """The 33 were never flagged already_present and are not a gap in the ledger's completeness."""
    claim = re.compile(r"flagged\s+`?already_present`?|ledger-completeness gap", re.I)
    for path in [README, AUDIT_PAGE, ROOT / "file_and_folder_structure.md"] + sorted((ROOT / "docs").glob("*.md")):
        text = _flat(path)
        hit = claim.search(text)
        assert not hit, (path.name, text[max(0, hit.start() - 120):hit.end() + 60])
